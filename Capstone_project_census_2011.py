import pandas as pd
from pymongo import MongoClient
from docx import Document
import mysql.connector
from mysql.connector import Error

# ---------------------------------------------------
# Data Pipeline and Analysis Script
# ---------------------------------------------------

# Helper function to truncate column names to a maximum length
def truncate_column_name(column_name, max_length=64):
    return column_name[:max_length]

# ---------------------------------------------------
# Task 1: Rename the Column Names
# ---------------------------------------------------

def rename_columns(df):
    """
    Renames specific columns in the DataFrame to have more appropriate names.
    """
    df.rename(columns={
        'State name': 'State/UT',
        'District name': 'District',
        'Male_Literate': 'Literate_Male',
        'Female_Literate': 'Literate_Female',
        'Rural_Households': 'Households_Rural',
        'Urban_Households': 'Households_Urban',
        'Age_Group_0_29': 'Young_and_Adult',
        'Age_Group_30_49': 'Middle_Aged',
        'Age_Group_50': 'Senior_Citizen',
        'Age not stated': 'Age_Not_Stated'
    }, inplace=True)
    return df

# ---------------------------------------------------
# Task 2: Rename State/UT Names
# ---------------------------------------------------

def format_state_ut_name(name):
    """
    Capitalizes the first letter of each word in the State/UT name, except for 'and'.
    """
    words = name.split()
    formatted_words = [word.capitalize() if word.lower() != 'and' else 'and' for word in words]
    return ' '.join(formatted_words)

def rename_state_ut_names(df):
    """
    Applies formatting to the 'State/UT' column for consistency.
    """
    df['State/UT'] = df['State/UT'].apply(format_state_ut_name)
    return df

# ---------------------------------------------------
# Task 3: Handle New State/UT Formation
# ---------------------------------------------------

def handle_state_ut_formation(df):
    """
    Updates the 'State/UT' column for districts affected by the formation of Telangana in 2014 and Ladakh in 2019.
    """
    # Telangana formation in 2014
    doc = Document('C:/Users/AJAY/Downloads/Telangana.docx')
    telangana_districts = []
    
    for para in doc.paragraphs:
        telangana_districts.extend([line.strip() for line in para.text.split('\n') if line.strip()])

    # Ensure specific districts are included
    telangana_districts.extend([
        "Adilabad", "Nizamabad", "Karimnagar", "Medak", 
        "Hyderabad", "Rangareddy", "Mahbubnagar", 
        "Nalgonda", "Warangal", "Khammam"
    ])
    
    # Remove duplicates if any
    telangana_districts = list(set(telangana_districts))

    # Update the DataFrame for Telangana
    df.loc[(df['District'].isin(telangana_districts)) & (df['State/UT'] == 'Andhra Pradesh'), 'State/UT'] = 'Telangana'

    # Ladakh formation in 2019
    ladakh_districts = ["Leh", "Kargil"]
    df.loc[(df['District'].isin(ladakh_districts)) & (df['State/UT'] == 'Jammu and Kashmir'), 'State/UT'] = 'Ladakh'
    
    return df

# ---------------------------------------------------
# Task 4: Find and Process Missing Data
# ---------------------------------------------------

def fill_missing_data(df):
    """
    Fills missing data for Population, Literate, and Households columns, and calculates the reduction in missing data.
    """
    # Calculate percentage of missing data before filling
    missing_before = df.isnull().mean() * 100
    
    # Fill missing values in Population, Literate, and Households columns
    if 'Population' in df.columns and 'Male' in df.columns and 'Female' in df.columns:
        df['Population'] = df['Population'].fillna(df['Male'] + df['Female'])
    
    if 'Literate' in df.columns and 'Literate_Male' in df.columns and 'Literate_Female' in df.columns:
        df['Literate'] = df['Literate'].fillna(df['Literate_Male'] + df['Literate_Female'])
    
    if 'Young_and_Adult' in df.columns and 'Middle_Aged' in df.columns and 'Senior_Citizen' in df.columns and 'Age_Not_Stated' in df.columns:
        df['Population'] = df['Population'].fillna(
            df['Young_and_Adult'] + df['Middle_Aged'] + df['Senior_Citizen'] + df['Age_Not_Stated']
        )
    
    if 'Households' in df.columns and 'Households_Rural' in df.columns and 'Households_Urban' in df.columns:
        df['Households'] = df['Households'].fillna(df['Households_Rural'] + df['Households_Urban'])

    # Calculate percentage of missing data after filling
    missing_after = df.isnull().mean() * 100

    # Create a DataFrame to compare missing data before and after
    comparison = pd.DataFrame({
        'Missing Before (%)': missing_before,
        'Missing After (%)': missing_after,
        'Reduction (%)': missing_before - missing_after
    })

    print("Missing Data Comparison:")
    print(comparison)

    return df

# ---------------------------------------------------
# Task 5: Save Data to MongoDB
# ---------------------------------------------------

def save_to_mongodb(df):
    """
    Saves the DataFrame to a MongoDB collection named 'census'.
    """
    client = MongoClient('mongodb://localhost:27017/')
    db = client['census_db']
    collection = db['census']  # Collection name 'census'

    # Convert DataFrame to a dictionary and insert it into MongoDB
    data_dict = df.to_dict("records")
    collection.insert_many(data_dict)

    print("Data has been successfully saved to MongoDB.")

# ---------------------------------------------------
# Task 6: Fetch Data from MongoDB and Save to MySQL
# ---------------------------------------------------

def fetch_from_mongodb_and_save_to_sql():
    """
    Fetches data from MongoDB, truncates column names, and saves it to a MySQL table.
    """
    try:
        # Connect to MongoDB
        client = MongoClient('mongodb://localhost:27017/')
        db = client['census_db']
        collection = db['census']

        # Fetch the data from MongoDB
        data = pd.DataFrame(list(collection.find()))

        # Remove the `_id` field
        if '_id' in data.columns:
            data = data.drop(columns=['_id'])

        # Print DataFrame columns for debugging
        print("Original DataFrame Columns:")
        print(data.columns.tolist())

        # Truncate long column names
        truncated_columns = {col: truncate_column_name(col) for col in data.columns}
        data.rename(columns=truncated_columns, inplace=True)

        # Print truncated DataFrame columns for debugging
        print("Truncated DataFrame Columns:")
        print(data.columns.tolist())

        # Connect to the relational database (MySQL)
        connection = mysql.connector.connect(
            user='root',
            password='raju10ajay11', 
            host='localhost',
            database='census'
        )

        cursor = connection.cursor()

        # Check and update the table schema
        cursor.execute("DESCRIBE census_2011;")
        existing_columns = [row[0] for row in cursor.fetchall()]

        missing_columns = [col for col in data.columns if col not in existing_columns]
        for col in missing_columns:
            cursor.execute(f"ALTER TABLE census_2011 ADD COLUMN `{col}` TEXT;")
            print(f"Added missing column `{col}` to table.")

        # Dynamically create the SQL table based on DataFrame columns
        table_name = 'census_2011'
        columns = data.columns.tolist()
        column_definitions = ', '.join([f'`{col}` TEXT' for col in columns])
        
        create_table_query = f"""
        CREATE TABLE IF NOT EXISTS `{table_name}` (
            id INT AUTO_INCREMENT PRIMARY KEY,
            {column_definitions}
        )
        """
        cursor.execute(create_table_query)

        # Print create table query for debugging
        print("Create Table Query:")
        print(create_table_query)

        # Insert data into the table
        for index, row in data.iterrows():
            placeholders = ', '.join(['%s'] * len(columns))
            insert_query = f"INSERT INTO `{table_name}` ({', '.join([f'`{col}`' for col in columns])}) VALUES ({placeholders})"
            cursor.execute(insert_query, tuple(row[col] for col in columns))

        # Commit the transaction
        connection.commit()

        print(f"Data has been successfully uploaded to the `{table_name}` table in the relational database.")

    except Error as e:
        print(f"Error: {e}")

    finally:
        # Close the cursor and connection
        if cursor:
            cursor.close()
        if connection:
            connection.close()

# ---------------------------------------------------
# Main Function to execute all tasks
# ---------------------------------------------------

def main():
    # Load your data
    file_path = 'C:/Users/AJAY/Downloads/census_2011.xlsx'
    df = pd.read_excel(file_path)

    # Execute all tasks
    df = rename_columns(df)
    df = rename_state_ut_names(df)
    df = handle_state_ut_formation(df)
    df = fill_missing_data(df)

    # Print the final DataFrame after all modifications
    print("Final DataFrame after all modifications:")
    print(df)

    save_to_mongodb(df)
    fetch_from_mongodb_and_save_to_sql()

# Run the main function
if __name__ == "__main__":
    main()
